"""Small helpers to fill the professor's .docx template by cloning its own
paragraphs and tables, so every heading, caption and table keeps the
template's formatting."""
import copy

from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from docx.table import Table


def body_items(doc):
    return list(doc.element.body)


def clone_after(anchor_el, proto_el):
    new = copy.deepcopy(proto_el)
    anchor_el.addnext(new)
    return new


def set_runs(p_el, parts, doc):
    """Replace the text of a cloned paragraph. `parts` is a list of
    (text, bold) pairs; the formatting of the first run in the prototype is
    reused, with bold switched on or off per part."""
    para = Paragraph(p_el, doc._body)
    runs = p_el.findall(qn("w:r"))
    proto_rpr = None
    if runs:
        rpr = runs[0].find(qn("w:rPr"))
        proto_rpr = copy.deepcopy(rpr) if rpr is not None else None
    for r in runs:
        p_el.remove(r)
    for tag in ("w:proofErr", "w:bookmarkStart", "w:bookmarkEnd", "w:hyperlink"):
        for x in p_el.findall(qn(tag)):
            p_el.remove(x)
    if isinstance(parts, str):
        parts = [(parts, None)]
    for text, bold in parts:
        run = para.add_run(text)
        if proto_rpr is not None:
            run._r.insert(0, copy.deepcopy(proto_rpr))
        if bold is not None:
            run.bold = bold
    return para


def page_break_before(p_el):
    ppr = p_el.find(qn("w:pPr"))
    if ppr is None:
        ppr = p_el.makeelement(qn("w:pPr"), {})
        p_el.insert(0, ppr)
    el = ppr.makeelement(qn("w:pageBreakBefore"), {})
    ppr.insert(0, el)


def keep_with_next(p_el):
    ppr = p_el.find(qn("w:pPr"))
    if ppr.find(qn("w:keepNext")) is None:
        ppr.insert(0, ppr.makeelement(qn("w:keepNext"), {}))


def add_picture_paragraph(anchor_el, proto_el, doc, image, width_in):
    """Centered paragraph holding one picture, inserted after anchor."""
    el = clone_after(anchor_el, proto_el)
    para = set_runs(el, "", doc)
    for r in el.findall(qn("w:r")):
        el.remove(r)
    para.add_run().add_picture(image, width=Inches(width_in))
    return el


def fill_table(tbl_el, doc, rows, bold_last=False):
    """Resize the template results table to len(rows) (header kept) by cloning
    its body rows, then write the values in."""
    table = Table(tbl_el, doc._body)
    trs = tbl_el.findall(qn("w:tr"))
    header, plain, special = trs[0], trs[1], trs[-1]
    for tr in trs[1:]:
        tbl_el.remove(tr)
    for i, values in enumerate(rows):
        proto = special if (bold_last and i == len(rows) - 1) else plain
        if i % 2 == 1 and len(trs) > 2 and not (bold_last and i == len(rows) - 1):
            proto = trs[2]  # template alternates row shading
        tr = copy.deepcopy(proto)
        tbl_el.append(tr)
    table = Table(tbl_el, doc._body)
    for r, values in enumerate(rows, start=1):
        for c, val in enumerate(values):
            cell = table.rows[r].cells[c]
            p = cell.paragraphs[0]
            set_runs(p._p, str(val), doc)
            for extra in cell.paragraphs[1:]:
                extra._p.getparent().remove(extra._p)
    return table


def set_header_cells(tbl_el, doc, headers):
    table = Table(tbl_el, doc._body)
    for c, h in enumerate(headers):
        set_runs(table.rows[0].cells[c].paragraphs[0]._p, h, doc)
