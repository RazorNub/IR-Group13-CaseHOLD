"""Builds deliverables/IR_Group13_CaseHOLD_Slides.pptx. Numbers come from results/."""
import json
import os

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = os.path.join(os.path.dirname(__file__), "..")
R = lambda *p: os.path.join(ROOT, *p)

NAVY = RGBColor(0x1B, 0x36, 0x5D)
TEAL = RGBColor(0x00, 0x80, 0x80)
GREY = RGBColor(0x59, 0x59, 0x59)
LIGHT = RGBColor(0xF4, 0xF6, 0xF8)
ROWALT = RGBColor(0xF9, 0xFA, 0xFB)
HILITE = RGBColor(0xE6, 0xF2, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x22, 0x22, 0x22)
FONT = "Calibri"

full = pd.read_csv(R("results/test_full_pool.csv")).set_index("model")
five = pd.read_csv(R("results/test_five_choice.csv")).set_index("model")
cfg = json.load(open(R("results/chosen_settings.json")))
ix = json.load(open(R("results/index_stats.json")))
err = json.load(open(R("results/error_analysis.json")))
split = json.load(open(R("results/split_stats.json")))
val = pd.read_csv(R("results/validation_tuning.csv"))
H3, HFT, FT = "Hybrid RRF (TF-IDF+BM25+Dense)", "Hybrid RRF (BM25+Dense fine-tuned)", "Dense (fine-tuned)"
HK = f"Hybrid RRF (TF-IDF+BM25+Dense, k={cfg['rrf_k']})"
CE = "Hybrid + cross-encoder re-rank"


def f(m, k, t=full):
    return f"{t.loc[m, k]:.3f}"


def vmrr(setting):
    return f"{val[val['setting'] == setting]['MRR'].iloc[0]:.3f}"


tst = split["test"]
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
TOTAL = 17
count = [0]


def text(slide, x, y, w, h, paras, size=18, color=BLACK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space=6):
    """paras: list of paragraphs; each is a string or a list of (text, bold, color) runs."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        runs = [(para, False, color)] if isinstance(para, str) else para
        for t, bold, col in [(r + (color,))[:3] if len(r) == 2 else r for r in runs]:
            run = p.add_run()
            run.text = t
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.name = FONT
            run.font.color.rgb = col or color
    return tb


def rect(slide, x, y, w, h, fill, line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    else:
        s.line.fill.background()
    style = s._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style")
    if style is not None:
        s._element.remove(style)
    return s


def new_slide(title, notes, kicker=None):
    count[0] += 1
    s = prs.slides.add_slide(BLANK)
    text(s, 0.6, 0.35, 12.1, 0.8, [[(title, True, NAVY)]], size=30, anchor=MSO_ANCHOR.BOTTOM)
    rect(s, 0.6, 1.2, 12.1, 0.03, NAVY)
    if kicker:
        text(s, 0.6, 0.12, 12, 0.35, [[(kicker, True, TEAL)]], size=12)
    text(s, 0.6, 7.0, 9, 0.35, [[("Group 13 · Legal Retrieval using CaseHOLD · CSE4144 Information "
                                   "Retrieval", False, GREY)]], size=10)
    text(s, 11.7, 7.0, 1.0, 0.35, [[(f"{count[0]} / {TOTAL}", False, GREY)]], size=10, align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = notes
    return s


def table(slide, x, y, w, rows, col_w=None, size=14, highlight_last=False, row_h=0.42):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w),
                                   Inches(row_h * len(rows)))
    t = shape.table
    tbl_pr = shape._element.graphic.graphicData.tbl.tblPr
    tbl_pr.set("bandRow", "0")
    tbl_pr.set("firstRow", "0")
    style = tbl_pr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}tableStyleId")
    if style is not None:
        style.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # "No Style, Table Grid"
    if col_w:
        for i, cw in enumerate(col_w):
            t.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        t.rows[r].height = Inches(row_h)
        for c, val_ in enumerate(row):
            cell = t.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            last = highlight_last and r == len(rows) - 1
            cell.fill.fore_color.rgb = NAVY if r == 0 else HILITE if last else ROWALT if r % 2 else WHITE
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
            run = p.add_run()
            run.text = str(val_)
            run.font.size = Pt(size)
            run.font.name = FONT
            run.font.bold = r == 0 or last
            run.font.color.rgb = WHITE if r == 0 else BLACK
    return t


def picture(slide, path, x, y, w=None, h=None):
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    return slide.shapes.add_picture(R("figures", path), Inches(x), Inches(y), **kw)


# 1 ------------------------------------------------------------------ title
count[0] += 1
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, 0.25, 7.5, NAVY)
text(s, 1.0, 1.5, 11, 0.5, [[("CSE4144 Information Retrieval · Project 13", True, TEAL)]], size=18)
text(s, 1.0, 2.05, 11.3, 1.9, [[("Legal Retrieval using CaseHOLD", True, NAVY)]], size=44)
text(s, 1.0, 3.15, 11.3, 1.0, [[("Comparing TF-IDF, BM25, a Sentence Transformer and Hybrid RRF for finding "
                                  "the holding a court cites", False, GREY)]], size=22)
rect(s, 1.0, 4.35, 5.5, 0.03, NAVY)
text(s, 1.0, 4.6, 11.5, 1.6, [[("Group 13", True, BLACK)],
                              [("Vaibhav Bansal (23FE10CSE00333)", False, BLACK),
                               ("    VAIBHAV.23FE10CSE00333@muj.manipal.edu", False, GREY)],
                              [("Pranav Joshi (23FE10CSE00701)", False, BLACK),
                               ("    Pranav.23FE10CSE00701@muj.manipal.edu", False, GREY)],
                              "Ishaan Manocha (23FE10CSE00143)"], size=18, space=2)
text(s, 1.0, 6.55, 11, 0.4, [[("B.Tech CSE, Semester VII · Manipal University Jaipur · 2026-27",
                               False, GREY)]], size=13)
s.notes_slide.notes_text_frame.text = (
    "Presenter: Vaibhav.\n"
    "Good morning. We are Group 13 and our assigned project is Legal Retrieval using CaseHOLD. "
    "In short, we built a search system that takes a passage from a court judgment and finds the holding of "
    "the case it cites. We compared the four ranking methods from the brief: TF-IDF, BM25, a Sentence "
    "Transformer and a hybrid using Reciprocal Rank Fusion, and we tried two small extensions on top. "
    "I will cover the problem and the data, Pranav will explain the method and main results, and Ishaan "
    "will cover the extra experiments, error analysis, the demo and the conclusion.")

# 2 ------------------------------------------------------------------ problem
s = new_slide("The problem: finding the holding a court relies on",
              "Presenter: Vaibhav.\n"
              "A holding is the legal rule a case decided. US judgments cite earlier cases constantly, and after "
              "the citation there is often a short bracket that starts with 'holding that ...'. "
              "A lawyer writing an argument has the reverse problem: they know the point they want to make and "
              "they need the case whose holding supports it. So we treat it as search. The query is the "
              "passage that does the citing, and the documents are holding statements. "
              "The difficulty is that the passage and the holding often describe the same idea in different "
              "words, for example 'fears about a gas pipeline' versus 'stigma damages', and legal text has many "
              "rare names and citation numbers.", kicker="Introduction")
text(s, 0.6, 1.55, 6.3, 4.8, [
    "Courts back every legal claim with a citation to an earlier case.",
    "The citation is often followed by its holding: what that case actually decided.",
    "A lawyer needs the reverse: given the argument, find the case whose holding supports it.",
    [("Our task: ", True, NAVY), ("query = the citing passage, documents = holding statements, "
                                  "goal = rank the cited holding first.", False, None)],
], size=20, space=14)
rect(s, 7.3, 1.6, 5.4, 4.6, LIGHT)
rect(s, 7.3, 1.6, 0.08, 4.6, NAVY)
text(s, 7.6, 1.75, 5.0, 4.4, [
    [("Query (part of the citing passage)", True, TEAL)],
    "“...the anticipation of increased interference or potential hazards on the land remaining "
    "after a taking, which one court has labeled ‘stigma damages,’ can be considered in "
    "assessing severance damages. See, e.g., Vector Pipeline v. 68.55 Acres of Land (<HOLDING>)”",
    [("Document we should return", True, TEAL)],
    "holding that stigma damages were properly included in calculation of severance damages where "
    "proximity of gas pipeline could trigger fears about possible mishaps",
], size=17, space=10)

# 3 ------------------------------------------------------------------ dataset
s = new_slide("Dataset: CaseHOLD",
              "Presenter: Vaibhav.\n"
              "CaseHOLD was released by Zheng and others at ICAIL 2021. It has about 53,000 questions built "
              "from the Harvard Law Library case law corpus. Each question is a citing passage with the "
              "holding replaced by a <HOLDING> marker, plus five candidate holdings, one correct. "
              "An important detail: the four wrong answers were chosen as the holdings most similar to the "
              "correct one by TF-IDF, so they are on purpose hard to tell apart by word overlap. "
              "We used the version in the LexGLUE benchmark on Hugging Face, which has fixed train, validation "
              "and test splits; the numbers are in the table. Queries are long, about 136 words, and holdings "
              "short, about 25 words, so only a few words overlap.", kicker="About the dataset")
rows = [["Split", "Queries", "Unique holdings", "Words / query", "Words / holding"]]
for name, key in [("Train", "train"), ("Validation", "validation"), ("Test", "test")]:
    d = split[key]
    rows.append([name, f"{d['queries']:,}", f"{d['holdings']:,}", d["ctx_words"], d["hold_words"]])
table(s, 0.6, 1.6, 6.4, rows, col_w=[1.4, 1.1, 1.5, 1.2, 1.2], size=14)
text(s, 0.6, 3.5, 6.4, 3.2, [
    "From US federal and state court decisions (Harvard Law Library corpus).",
    "Each question: citing passage + 5 holdings, exactly one correct.",
    [("The 4 wrong holdings were picked by TF-IDF similarity ", True, NAVY),
     ("to the right one, to make the task hard.", False, None)],
    "We use the LexGLUE copy (coastalcph/lex_glue, case_hold).",
], size=18, space=12)
picture(s, "dataset_lengths.png", 7.1, 1.6, w=5.7)
text(s, 7.1, 3.98, 5.7, 0.5, [[("Word counts in the test split: long queries, short holdings", False, GREY)]],
     size=13, align=PP_ALIGN.CENTER)

# 4 ------------------------------------------------------------------ example
s = new_slide("One test question, as the models see it",
              "Presenter: Vaibhav.\n"
              "This is test question 0. The passage is about stigma damages near a gas pipeline. "
              "Option 4 is correct. Notice that options 1, 2 and 3 are also about severance or stigma damages. "
              "That is the TF-IDF based distractor selection I mentioned. A model that only counts shared words "
              "can easily be confused between 2, 3 and 4. In our full search setting the model does not even "
              "get these five options; it has to find option 4 among 13,297 holdings.", kicker="About the dataset")
rect(s, 0.6, 1.55, 12.1, 1.55, LIGHT)
rect(s, 0.6, 1.55, 0.08, 1.55, NAVY)
text(s, 0.85, 1.6, 11.7, 1.5, [[("Citing passage (end): ", True, NAVY),
                                ("…have held that the anticipation of increased interference or potential "
                                 "hazards on the land remaining after a taking, which one court has labeled "
                                 "“stigma damages,” can be considered in assessing severance damages. "
                                 "See, e.g., Vector Pipeline, L.P., v. 68.55 Acres of Land, 157 F.Supp.2d 949, 957 "
                                 "(N.D.Ill.2001) (<HOLDING>)", False, None)]], size=16)
opts = [("0", "holding that actual damages under respa included emotional damages", False),
        ("1", "holding that increased fear of flooding was a matter that factfinder could properly consider "
              "in assessing severance damages", False),
        ("2", "holding that local buyers fear of possible knapweed infestation legitimately affected "
              "calculation of severance damages", False),
        ("3", "recognizing concept of stigma damages as they apply to property which is in fact contaminated",
         False),
        ("4", "holding that stigma damages were properly included in calculation of severance damages where "
              "proximity of gas pipeline could trigger fears about possible mishaps", True)]
y = 3.3
for k, t_, ok in opts:
    if ok:
        rect(s, 0.6, y - 0.03, 12.1, 0.68, HILITE)
    text(s, 0.75, y, 0.6, 0.6, [[(k, True, TEAL if ok else NAVY)]], size=18)
    text(s, 1.3, y, 10.0, 0.65, [[(t_, ok, None)]], size=15)
    if ok:
        text(s, 11.3, y, 1.3, 0.6, [[("correct", True, TEAL)]], size=15, align=PP_ALIGN.RIGHT)
    y += 0.72

# 5 ------------------------------------------------------------------ setup
s = new_slide("Turning a multiple-choice dataset into a search task",
              "Presenter: Vaibhav.\n"
              "CaseHOLD is normally used as a multiple-choice reading test. For an IR project we need a "
              "document collection, so we pooled every distinct holding from the test split: 13,297 documents. "
              "Every one of the 3,600 test queries has exactly one relevant document in that pool. That is our "
              "main, realistic setting. We also report the original five-choice setting, where each query is "
              "only compared with its own five options; there random guessing gives 20 percent. "
              "One more detail: 1,032 correct holdings also show up as a wrong option for other questions, so "
              "the same text can be relevant for one query and irrelevant for another. "
              "Important for honesty: we chose all settings on the validation split and only ran the test split "
              "once at the end. Pranav will now explain the method.", kicker="Problem setup")
for i, (head, body, x) in enumerate([
        ("Full-pool setting (main)",
         [f"Documents: all {tst['holdings']:,} unique holdings of the test split",
          f"Queries: {tst['queries']:,} citing passages",
          "Exactly 1 relevant holding per query",
          "Metrics: Precision@10, Recall@k, MAP, nDCG@10"], 0.6),
        ("5-choice setting (original task)",
         ["Each query ranks only its own 5 options",
          "Accuracy = correct option ranked first",
          "Random guessing = 0.20",
          "Hard: options were chosen to look alike"], 6.85)]):
    rect(s, x, 1.6, 5.85, 3.6, LIGHT)
    rect(s, x, 1.6, 5.85, 0.08, NAVY if i == 0 else TEAL)
    text(s, x + 0.3, 1.8, 5.3, 0.6, [[(head, True, NAVY if i == 0 else TEAL)]], size=22)
    text(s, x + 0.3, 2.5, 5.3, 2.6, body, size=17, space=10)
text(s, 0.6, 5.5, 12.1, 1.2, [
    [("Protocol: ", True, NAVY), ("every parameter was tuned on the 3,900 validation queries; the test split "
                                  "was used once, at the end. The train split was used only for the "
                                  "fine-tuning extension.", False, None)]], size=17)

# 6 ------------------------------------------------------------------ pipeline
s = new_slide("System pipeline",
              "Presenter: Pranav.\n"
              "This is the whole system. The query and the holdings go down two paths. "
              "On the top, the sparse path: we clean and stem the text, build an inverted index, and score with "
              "TF-IDF and BM25. On the bottom, the dense path: the raw text goes into a small Sentence "
              "Transformer called all-MiniLM-L6-v2, which turns each text into a 384-number vector, and we "
              "compare vectors with cosine similarity. The ranked lists are merged with Reciprocal Rank Fusion. "
              "The cross-encoder box on the right is an extra step we tried; I'll leave it to Ishaan.",
              kicker="Methodology")
picture(s, "pipeline.png", 0.9, 1.55, w=11.5)

# 7 ------------------------------------------------------------------ preprocessing
s = new_slide("Preprocessing and the inverted index",
              "Presenter: Pranav.\n"
              "For the sparse models we lowercase, remove the HOLDING marker, and split into words. "
              "We drop any token with a digit, because citations like '157 F.Supp.2d 949' are unique to one "
              "case and only add noise. We remove English stop words plus a few legal words that appear "
              "everywhere: 'holding' starts 91 percent of all holdings, so it tells us nothing. Then the Porter "
              "stemmer, so 'damages' and 'damage' match. On validation, stemming clearly helped TF-IDF, while stop "
              "words alone did not. "
              "The inverted index maps each stemmed term to a postings list of holding ids and term counts. "
              "We wrote BM25 ourselves on top of it, and checked that our fast matrix version gives the same "
              "scores as a plain loop over the postings lists. The dense model does not use any of this; it "
              "reads the original sentences.", kicker="Methodology")
steps = [("1", "Lowercase, remove <HOLDING>"), ("2", "Tokenise; drop tokens with digits (citations)"),
         ("3", "Remove stop words + legal fillers (holding, v, see, id…)"), ("4", "Porter stemming")]
y = 1.65
for n, t_ in steps:
    rect(s, 0.6, y, 0.55, 0.55, NAVY)
    text(s, 0.6, y, 0.55, 0.55, [[(n, True, WHITE)]], size=18, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 1.3, y, 5.4, 0.6, [t_], size=18, anchor=MSO_ANCHOR.MIDDLE)
    y += 0.75
rect(s, 0.6, 4.85, 6.1, 1.45, LIGHT)
text(s, 0.8, 4.95, 5.8, 1.7, [
    [("Holding: ", True, NAVY), ("“holding that stigma damages were properly included in calculation of "
                                 "severance damages”", False, None)],
    [("Tokens: ", True, TEAL), ("stigma, damag, properli, includ, calcul, sever, damag", False, None)]],
    size=15, space=8)
text(s, 7.2, 1.55, 5.5, 0.5, [[("Inverted index on the test pool", True, NAVY)]], size=20)
table(s, 7.2, 2.15, 5.5, [["", "Value"], ["Holdings indexed", f"{ix['documents']:,}"],
                          ["Vocabulary (stemmed terms)", f"{ix['vocabulary']:,}"],
                          ["Postings (term, holding) pairs", f"{ix['postings']:,}"],
                          ["Avg. terms per holding", ix["avg_doc_length"]],
                          ["Avg. postings list length", ix["avg_postings_length"]]],
      col_w=[3.6, 1.9], size=15)
text(s, 7.2, 4.85, 5.5, 1.9, [
    [("Validation check (TF-IDF MRR)", True, NAVY)],
    f"no stop words, no stemming: {vmrr('TF-IDF stopwords=False stemming=False')}",
    f"stop words only: {vmrr('TF-IDF stopwords=True stemming=False')}",
    f"stop words + stemming: {vmrr('TF-IDF stopwords=True stemming=True')}"], size=15, space=4)

# 8 ------------------------------------------------------------------ models
s = new_slide("The four ranking models",
              "Presenter: Pranav.\n"
              "TF-IDF: each text is a vector of term weights, log term frequency times inverse document "
              "frequency, and we rank by cosine similarity. "
              "BM25 is the probabilistic model. k1 controls how fast repeated terms stop adding score, and b "
              "controls how much long documents are penalised. "
              "The Sentence Transformer is all-MiniLM-L6-v2, a 6-layer model, about 90 MB, so it runs on a "
              "laptop. It reads up to 256 word pieces. "
              "RRF: for each document we add 1 over k plus its rank in each list. It only uses ranks, which is "
              "useful because BM25 scores are around 40 while cosine scores are between minus 1 and 1, so we "
              "cannot add them directly. k = 60 is the standard value from the original paper by Cormack and "
              "others.", kicker="Methodology")
cards = [("TF-IDF", "Vector of (1 + log tf) × idf weights, ranked by cosine similarity."),
         ("BM25", "Probabilistic model with term saturation (k1) and length normalisation (b). "
                  "Our own implementation on the inverted index."),
         ("Sentence Transformer", "all-MiniLM-L6-v2: text → 384-d vector; cosine similarity. "
                                  "Captures meaning, not exact words."),
         ("Hybrid RRF", "Merges the top-1000 lists of the three models by rank only, so no score "
                        "normalisation is needed.")]
for i, (h, b) in enumerate(cards):
    x = 0.6 + i * 3.08
    rect(s, x, 1.6, 2.88, 2.55, LIGHT)
    rect(s, x, 1.6, 2.88, 0.08, NAVY if i < 3 else TEAL)
    text(s, x + 0.15, 1.78, 2.6, 0.5, [[(h, True, NAVY if i < 3 else TEAL)]], size=19)
    text(s, x + 0.15, 2.35, 2.6, 1.8, [b], size=17)
text(s, 0.6, 4.45, 12.1, 0.5, [[("BM25(q, d) = Σ idf(t) · tf·(k1 + 1) / ( tf + k1·(1 - b + "
                                 "b·|d|/avgdl) )", True, NAVY)]], size=20, align=PP_ALIGN.CENTER)
text(s, 0.6, 5.15, 12.1, 0.5, [[("RRF(d) = Σ over models m of  1 / (k + rank_m(d)),   k = 60", True, TEAL)]],
     size=20, align=PP_ALIGN.CENTER)
text(s, 0.6, 5.9, 12.1, 0.8, [[("Metrics: ", True, NAVY),
                               ("one relevant holding per query, so MAP = MRR and the best possible "
                                "Precision@10 is 0.1. nDCG@10 = 1 / log2(rank + 1) inside the top 10.",
                                False, None)]], size=16, align=PP_ALIGN.CENTER)

# 9 ------------------------------------------------------------------ tuning
s = new_slide("Choosing settings on the validation split",
              "Presenter: Pranav.\n"
              f"We never touched the test split while choosing settings. On the 3,900 validation queries: "
              f"using the whole passage as the query was best; cutting it to the last 32 words dropped BM25 MRR "
              f"from {vmrr('BM25 last all words')} to {vmrr('BM25 last 32 words')}. "
              f"The heat map is the BM25 grid. The best was k1 = {cfg['bm25_k1']} and b = {cfg['bm25_b']}. A low "
              "k1 makes sense because holdings are short, so a repeated word should not count much more, and a "
              "high b normalises length strongly. "
              f"For the dense model, MiniLM beat BGE-small ({vmrr('all-MiniLM-L6-v2 last all words')} vs "
              f"{vmrr('bge-small-en-v1.5 last 96 words')}). "
              f"For RRF, smaller k was slightly better on validation, k = 1 gave {vmrr('k=1')} versus "
              f"{vmrr('k=60')} at k = 60. We kept 60 for the main results because it is the standard value, "
              "and we show k = 1 as an extra row.", kicker="Methodology")
picture(s, "bm25_tuning.png", 0.6, 1.5, h=4.6)
text(s, 7.0, 1.6, 5.7, 5.0, [
    [("Query: ", True, NAVY), (f"whole passage (MRR {vmrr('BM25 last all words')}) beat last 64 words "
                               f"({vmrr('BM25 last 64 words')}) and last 32 ({vmrr('BM25 last 32 words')})",
                               False, None)],
    [("BM25: ", True, NAVY), (f"k1 = {cfg['bm25_k1']}, b = {cfg['bm25_b']} (short documents: low k1, "
                              "strong length normalisation)", False, None)],
    [("Dense model: ", True, NAVY), (f"MiniLM {vmrr('all-MiniLM-L6-v2 last all words')} vs BGE-small "
                                     f"{vmrr('bge-small-en-v1.5 last 96 words')}", False, None)],
    [("RRF k: ", True, NAVY), (f"k = 1 → {vmrr('k=1')}, k = 60 → {vmrr('k=60')}; we report "
                               "both", False, None)]], size=18, space=16)

# 10 ----------------------------------------------------------------- full pool table
s = new_slide("Results: full-pool search (3,600 queries, 13,297 holdings)",
              "Presenter: Pranav.\n"
              f"These are the four models from the brief on the test split. BM25 is the best single model, MAP "
              f"{f('BM25', 'MAP')}, TF-IDF is close behind. The general-purpose dense model is clearly worse at "
              f"{f('Dense', 'MAP')}: it was trained on web text, not court judgments. "
              f"Remember Precision@10 can be at most 0.1 here because there is only one right answer. "
              f"The hybrid gives the best numbers of the four. The MAP gain is small, {f('BM25', 'MAP')} to "
              f"{f(H3, 'MAP')}, but recall at 100 goes from {f('BM25', 'R@100')} to {f(H3, 'R@100')}. So fusion "
              "mostly rescues answers that BM25 put deep in the list. A Wilcoxon signed-rank test says the "
              "hybrid's improvement over BM25 is significant. The k = 1 version, chosen on validation, is a "
              "bit better again.", kicker="Results")
rows = [["Ranking model", "Precision@10", "Recall@10", "Recall@100", "MAP (= MRR)", "nDCG@10"]]
for m, n in [("TF-IDF", "TF-IDF"), ("BM25", "BM25"), ("Dense", "Sentence Transformer"),
             (H3, "Hybrid RRF (k = 60, standard)"), (HK, f"Hybrid RRF (k = {cfg['rrf_k']}, tuned on validation)")]:
    rows.append([n, f(m, "P@10"), f(m, "R@10"), f(m, "R@100"), f(m, "MAP"), f(m, "nDCG@10")])
table(s, 0.6, 1.6, 12.1, rows, col_w=[3.6, 1.7, 1.7, 1.7, 1.7, 1.7], size=17, highlight_last=True, row_h=0.55)
text(s, 0.6, 5.2, 12.1, 1.6, [
    [("BM25 is the strongest single model; ", True, NAVY),
     ("the off-the-shelf dense model is clearly weaker on legal text.", False, None)],
    [("The hybrid mainly improves recall: ", True, NAVY),
     (f"Recall@100 {f('BM25', 'R@100')} → {f(H3, 'R@100')}, MAP {f('BM25', 'MAP')} → "
      f"{f(H3, 'MAP')} (Wilcoxon p < 0.001 vs BM25).", False, None)]], size=17, space=8)

# 11 ----------------------------------------------------------------- PR graph
s = new_slide("Precision-Recall graph and Recall@k",
              "Presenter: Pranav.\n"
              "The brief asks for a precision-recall graph. With one relevant document per query, the classic "
              "11-point curve would be flat, so we plot mean precision at k against mean recall at k for k from "
              "1 to 100. Precision is on a log scale since it falls roughly like 1 over k. "
              "The hybrid curve, in red, is level with BM25 at the first two cut-offs and above it from k = 3. "
              "The dense model is lowest everywhere. On the right is recall at k: the green and purple lines "
              "are the fine-tuned model and its hybrid, which Ishaan will explain next.", kicker="Results")
picture(s, "pr_curve.png", 0.6, 1.5, h=4.9)
picture(s, "recall_at_k.png", 6.9, 1.5, h=4.9)

# 12 ----------------------------------------------------------------- five choice
s = new_slide("Results: the original 5-choice task",
              "Presenter: Ishaan.\n"
              "In the original multiple-choice format every model is far above the 20 percent of random "
              f"guessing. BM25 gets {f('BM25', 'accuracy', five)}. Interesting point: here the hybrid "
              f"({f(H3, 'accuracy', five)}) is slightly below BM25. Why? The four wrong options were picked by "
              "TF-IDF similarity to the right one, so all five share vocabulary, and the extra lexical signal "
              "from fusion does not separate them. "
              f"The best is the fine-tuned dense model at {f(FT, 'accuracy', five)}. For context the CaseHOLD "
              "paper reports a macro F1 of about 0.4 for a BiLSTM and about 0.6 for fine-tuned BERT, on their "
              "own split and a different metric, so it is not a direct comparison, but our numbers are in a "
              "sensible range.", kicker="Results")
picture(s, "five_choice_accuracy.png", 0.6, 1.5, w=7.3)
text(s, 8.3, 1.7, 4.4, 4.8, [
    [("Random guess = 0.20", True, NAVY)],
    f"BM25 {f('BM25', 'accuracy', five)} beats the 3-way hybrid {f(H3, 'accuracy', five)} here.",
    "Reason: the wrong options were chosen by TF-IDF similarity, so all five look alike to word "
    "matching.",
    f"Best: fine-tuned dense model, {f(FT, 'accuracy', five)}.",
    [("Published reference: ", True, NAVY), ("BiLSTM about 0.4 and BERT about 0.6 macro F1 (Zheng et al., "
                                             "different split and metric).", False, None)]], size=17, space=12)

# 13 ----------------------------------------------------------------- extensions
s = new_slide("Extensions: fine-tuning helps, a web re-ranker hurts",
              "Presenter: Ishaan.\n"
              "Since the dense model was weak, we asked whether it just lacks legal knowledge. We fine-tuned the "
              "same MiniLM for one epoch on the 45,000 training questions: pairs of passage and correct holding, "
              "with one wrong option as a hard negative, using the multiple negatives ranking loss. That means "
              "inside a batch of 32 the model must pick the right holding over all others. It took about 11 "
              "minutes on the laptop GPU. "
              f"Its MAP went from {f('Dense', 'MAP')} to {f(FT, 'MAP')}, better than BM25, and fusing it with BM25 "
              f"gives our best result, {f(HFT, 'MAP')}, with recall at 100 of {f(HFT, 'R@100')}. "
              "The second extension was a cross-encoder re-ranker trained on MS MARCO web search. It made things "
              f"worse: {f(CE, 'MAP')}. It learned whether a web page answers a short question, which is very "
              "different from matching a 140-word legal passage to a holding. Lesson: domain data matters more "
              "than model type.", kicker="Beyond the brief")
rows = [["Run (full pool, test)", "Recall@10", "Recall@100", "MAP", "nDCG@10"]]
for m, n in [("BM25", "BM25 (reference)"), (H3, "Hybrid RRF (TF-IDF + BM25 + Dense)"),
             (CE, "Hybrid + ms-marco cross-encoder"), ("Dense", "Sentence Transformer (base)"),
             (FT, "Sentence Transformer, fine-tuned"), (HFT, "Hybrid RRF (BM25 + fine-tuned)")]:
    rows.append([n, f(m, "R@10"), f(m, "R@100"), f(m, "MAP"), f(m, "nDCG@10")])
table(s, 0.6, 1.6, 8.0, rows, col_w=[3.5, 1.15, 1.25, 0.9, 1.2], size=15, highlight_last=True, row_h=0.55)
text(s, 8.95, 1.6, 3.8, 5.2, [
    [("Fine-tuning", True, NAVY)],
    "45,000 train pairs, 1 hard negative each, multiple negatives ranking loss, 1 epoch, ~11 min.",
    [("Cross-encoder", True, NAVY)],
    "ms-marco-MiniLM re-scores the top 20. Trained on web search, so it does not transfer to legal "
    "text.",
    [("Takeaway", True, TEAL)],
    "Domain data mattered more than a bigger or fancier model."], size=16, space=8)

# 14 ----------------------------------------------------------------- error analysis
s = new_slide("Error analysis: where and why the models fail",
              "Presenter: Ishaan.\n"
              f"BM25 and the dense model fail on different queries. BM25 gets rank 1 where the dense model does not "
              f"in {err['bm25_top1_dense_not']} queries, the dense model wins in {err['dense_top1_bm25_not']}. "
              "That disagreement is the reason fusion helps at all. "
              "Three examples. Query 3185 is about money structuring offences: the rare word 'structuring' lets "
              "BM25 rank the answer first, while the dense model ranks it 129th. Query 296 is the opposite: no "
              "shared keywords, BM25 ranks the answer 1,694th, the dense model ranks it first because it "
              "understands the meaning. Query 2885 fails everywhere: the passage is about a strip search, but "
              "the citation was used for a procedural point about one court panel being bound by another. "
              "Nothing in the passage says that, so no ranker could find it. That is a limit of the data. "
              f"Also, in {err['neither_top10'] / tst['queries']:.0%} of queries neither model has the answer in "
              "the top 10, and longer holdings are easier because there are more words to match.",
              kicker="Analysis")
stats = [(str(err["bm25_top1_dense_not"]), "queries where only BM25 ranks the answer first"),
         (str(err["dense_top1_bm25_not"]), "queries where only the dense model does"),
         (f"{err['neither_top10'] / tst['queries']:.0%}", "of queries: neither model has it in the top 10")]
for i, (big, small) in enumerate(stats):
    y = 1.6 + i * 1.3
    text(s, 0.6, y, 1.9, 1.0, [[(big, True, NAVY)]], size=36, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 2.5, y, 2.9, 1.0, [small], size=15, anchor=MSO_ANCHOR.MIDDLE)
rows = [["Query", "BM25 rank", "Dense rank", "What happened"],
        ["3185", "1", "129", "Rare word “structuring” matched exactly"],
        ["296", "1,694", "1", "Same meaning, no shared keywords"],
        ["2885", "3,475", "6,823", "Passage is about a strip search; the cited holding is about court "
                                   "procedure"]]
table(s, 5.7, 1.6, 7.0, rows, col_w=[0.9, 1.1, 1.2, 3.8], size=14, row_h=0.75)
text(s, 0.6, 5.65, 12.1, 1.1, [
    [("Longer holdings are easier: ", True, NAVY),
     (f"BM25 MRR {err['mrr_by_gold_length']['0-15 words']['BM25']:.3f} for holdings of 15 words or less vs "
      f"{err['mrr_by_gold_length']['41-1000 words']['BM25']:.3f} above 40 words.", False, None)]], size=17)

# 15 ----------------------------------------------------------------- demo
s = new_slide("Search interface (live demo)",
              "Presenter: Ishaan.\n"
              "For the demonstration we built a Streamlit app. Run it with: streamlit run app/search_app.py. "
              "You can pick any CaseHOLD test query by number or type your own passage. The app shows the query "
              "terms after preprocessing, then the top results of TF-IDF, BM25, the dense model and the hybrid "
              "side by side, and marks the correct holding. For query 0, TF-IDF, BM25 and the hybrid rank it "
              "first and the dense model ranks it second. You can also switch to the fine-tuned dense model in "
              "the sidebar. Suggested live demo: query 0, then 3185 and 296 from the error analysis, then type "
              "a passage of your own.", kicker="Demo")
picture(s, "search_app.png", 0.6, 1.45, h=5.45)
text(s, 8.5, 1.6, 4.2, 5.2, [
    [("streamlit run app/search_app.py", True, NAVY)],
    "Pick a test query or type any passage.",
    "Shows the stemmed query terms.",
    "Top results of TF-IDF, BM25, dense and hybrid side by side.",
    "Marks the correct holding when it is known.",
    "Base or fine-tuned dense model from the sidebar."], size=16, space=10)

# 16 ----------------------------------------------------------------- conclusion
s = new_slide("Conclusion and limitations",
              "Presenter: Ishaan.\n"
              "To sum up. BM25 is a very strong baseline on legal text and beat the general-purpose dense "
              "model clearly. Reciprocal Rank Fusion helped mostly with recall, because the dense list on its "
              "own is weak. The biggest gain came from training on legal data: one epoch of fine-tuning made "
              "the small dense model better than BM25, and fusing the two was best overall. "
              "Limits: one relevant holding per query, so a holding that is legally fine but not the one cited "
              "counts as wrong, and Precision@10 is capped at 0.1. Our pool is the test split's holdings, not a "
              "whole case law database. And we used small models because everything ran on one laptop. "
              "Next we would try a legal model such as Legal-BERT, train the re-ranker on CaseHOLD, and give "
              "the rankers different weights inside RRF. Thank you, we are happy to take questions.",
              kicker="Conclusion")
text(s, 0.6, 1.55, 6.0, 0.5, [[("What we found", True, NAVY)]], size=22)
text(s, 0.6, 2.15, 6.0, 4.5, [
    f"BM25 (MAP {f('BM25', 'MAP')}) beat the off-the-shelf dense model ({f('Dense', 'MAP')}).",
    f"RRF mainly raised recall: Recall@100 {f('BM25', 'R@100')} → {f(H3, 'R@100')}.",
    f"Fine-tuning on CaseHOLD: dense {f('Dense', 'MAP')} → {f(FT, 'MAP')}; with BM25 "
    f"{f(HFT, 'MAP')}, our best.",
    "A web-trained cross-encoder made results worse."], size=18, space=14)
rect(s, 6.9, 1.55, 5.8, 5.1, LIGHT)
text(s, 7.15, 1.7, 5.4, 0.5, [[("Limitations", True, TEAL)]], size=22)
text(s, 7.15, 2.3, 5.4, 2.2, [
    "One relevant holding per query; Precision@10 capped at 0.1.",
    "Pool = test-split holdings, not a full case law database.",
    "Small models, 1 training epoch, one laptop."], size=16, space=8)
text(s, 7.15, 4.35, 5.4, 0.5, [[("Next steps", True, TEAL)]], size=22)
text(s, 7.15, 4.95, 5.4, 1.6, [
    "Legal-domain encoder (e.g. Legal-BERT).",
    "Train the re-ranker on CaseHOLD itself.",
    "Weighted RRF instead of equal weights."], size=16, space=8)

# 17 ----------------------------------------------------------------- references
s = new_slide("References",
              "Presenter: Ishaan.\nKeep this slide up during questions. Code, results and figures are in the "
              "project folder; run_all.sh reproduces every number.", kicker="References")
text(s, 0.6, 1.55, 12.1, 5.2, [
    "Zheng, Guha, Anderson, Henderson, Ho. When Does Pretraining Help? Assessing Self-Supervised Learning "
    "for Law and the CaseHOLD Dataset. ICAIL 2021.",
    "Chalkidis et al. LexGLUE: A Benchmark Dataset for Legal Language Understanding in English. ACL 2022.",
    "Robertson, Zaragoza. The Probabilistic Relevance Framework: BM25 and Beyond. FnTIR 3(4), 2009.",
    "Cormack, Clarke, Buettcher. Reciprocal Rank Fusion outperforms Condorcet and individual Rank Learning "
    "Methods. SIGIR 2009.",
    "Reimers, Gurevych. Sentence-BERT: Sentence Embeddings using Siamese BERT-Networks. EMNLP 2019.",
    "Wang et al. MiniLM: Deep Self-Attention Distillation. NeurIPS 2020.",
    "Henderson et al. Efficient Natural Language Response Suggestion for Smart Reply. arXiv 2017.",
    "Nogueira, Cho. Passage Re-ranking with BERT. arXiv 2019.",
    "Manning, Raghavan, Schütze. Introduction to Information Retrieval. Cambridge University Press, 2008.",
], size=15, space=9)

assert count[0] == TOTAL, count[0]
out = R("deliverables/IR_Group13_CaseHOLD_Slides.pptx")
prs.save(out)
print("saved", out)
