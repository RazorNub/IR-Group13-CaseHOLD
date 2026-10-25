"""Builds the viva preparation sheet (deliverables/IR_Group13_Viva_QA.docx)."""
import json
import os

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = os.path.join(os.path.dirname(__file__), "..")
R = lambda *p: os.path.join(ROOT, *p)
NAVY, TEAL = RGBColor(0x1B, 0x36, 0x5D), RGBColor(0x00, 0x80, 0x80)

full = pd.read_csv(R("results/test_full_pool.csv")).set_index("model")
five = pd.read_csv(R("results/test_five_choice.csv")).set_index("model")
cfg = json.load(open(R("results/chosen_settings.json")))
ix = json.load(open(R("results/index_stats.json")))
err = json.load(open(R("results/error_analysis.json")))
split = json.load(open(R("results/split_stats.json")))
H3, HFT, FT = "Hybrid RRF (TF-IDF+BM25+Dense)", "Hybrid RRF (BM25+Dense fine-tuned)", "Dense (fine-tuned)"
HK = f"Hybrid RRF (TF-IDF+BM25+Dense, k={cfg['rrf_k']})"
CE = "Hybrid + cross-encoder re-rank"
f = lambda m, k, t=full: f"{t.loc[m, k]:.3f}"
tst = split["test"]

QA = [
    ("About the task and data", [
        ("What is CaseHOLD?",
         "A dataset of about 53,000 questions built from US court decisions (Harvard Law Library corpus), released by "
         "Zheng et al. at ICAIL 2021. Each question is a passage that cites a case, with the holding replaced by "
         "<HOLDING>, plus five candidate holdings of which one is correct."),
        ("What is a holding?",
         "The legal rule or decision a case established. In judgments it often appears in brackets after a citation, "
         "starting with 'holding that ...'."),
        ("CaseHOLD is a multiple-choice dataset. How is this an IR project?",
         f"We pooled all {tst['holdings']:,} distinct holdings in the test split into one collection. Each of the "
         f"{tst['queries']:,} citing passages is a query with exactly one relevant holding, and every model has to "
         "rank the whole collection. We also report the original 5-choice accuracy as a second setting."),
        ("Which version of the data did you use, and why that one?",
         "The LexGLUE copy on Hugging Face (coastalcph/lex_glue, config case_hold): 45,000 train, 3,900 validation, "
         "3,600 test. It loads directly with the datasets library and has fixed splits, so our numbers are "
         "reproducible."),
        ("How were the wrong options chosen?",
         "In the original paper, by TF-IDF similarity to the correct holding (the most similar ones, below a 0.75 "
         "cut-off). That is why they look alike and why the 5-choice task is hard for word matching."),
        ("Did you use a subset?",
         "No. All 3,600 test queries and all 13,297 test holdings were used. Fine-tuning used the full 45,000 "
         "training examples for one epoch."),
    ]),
    ("Method", [
        ("Walk us through your preprocessing.",
         "Lowercase, remove the <HOLDING> marker, join dotted abbreviations, split on non-alphanumerics, drop "
         "one-letter tokens and any token with a digit (citations like 157 F.Supp.2d 949), remove English stop "
         "words plus legal fillers such as 'holding', 'v', 'see', 'id', then Porter stemming. The dense model gets "
         "the raw text instead."),
        ("Why remove tokens with digits?",
         "They are mostly reporter citations and page numbers. They are unique per case, so they never help match "
         "a passage to a holding and only add noise and index size."),
        ("Why did you add 'holding' as a stop word?",
         "91% of holdings start with 'holding that'. A word in almost every document has no power to tell "
         "documents apart."),
        ("What exactly is in your inverted index?",
         f"For each stemmed term, a postings list of (holding id, term frequency) pairs, plus each holding's length. "
         f"Test pool: {ix['vocabulary']:,} terms, {ix['postings']:,} postings, {ix['avg_doc_length']} terms per "
         "holding on average."),
        ("Did you use a library for BM25?",
         "No, we implemented it on top of our inverted index (src/index.py). For speed we also turn the postings "
         "into a sparse term-by-document weight matrix, and we checked that both versions give the same scores "
         "(differences below 1e-5)."),
        ("Explain BM25's k1 and b.",
         "k1 controls term frequency saturation: after a few occurrences extra repeats add little. b controls "
         "length normalisation: b = 1 fully normalises by document length, b = 0 ignores length. We picked "
         f"k1 = {cfg['bm25_k1']}, b = {cfg['bm25_b']} on validation; short holdings rarely repeat words, so a low k1 "
         "fits."),
        ("How is TF-IDF different from BM25 here?",
         "TF-IDF uses (1 + log tf) x idf weights and cosine similarity between the query and document vectors. "
         "BM25 adds the saturation and length normalisation from a probabilistic model. On our data the two are "
         f"close (MAP {f('TF-IDF', 'MAP')} vs {f('BM25', 'MAP')})."),
        ("Which Sentence Transformer did you use and why?",
         "all-MiniLM-L6-v2: 6 layers, 384-dimensional vectors, small enough for a laptop. On validation it beat "
         "bge-small-en-v1.5, the other small model we tried."),
        ("What is Reciprocal Rank Fusion?",
         "For every document, add 1 / (k + rank) over each model's ranked list. It uses only rank positions, so we "
         "don't have to make BM25 scores (around 40) comparable with cosine scores (-1 to 1). k = 60 is the "
         "standard value from Cormack et al. 2009."),
        ("Why did you report both k = 60 and k = 1?",
         f"k = 60 is the standard value the brief mentions, so it is our main result. On validation k = 1 was "
         f"slightly better, so we report it as an extra row (test MAP {f(HK, 'MAP')} vs {f(H3, 'MAP')}). Small k "
         "gives more weight to the very top of each list."),
        ("How did you avoid overfitting to the test set?",
         "Every choice (query length, k1, b, preprocessing, dense model, RRF k) was made on the validation split. "
         "The test split was run once at the end."),
    ]),
    ("Metrics", [
        ("Why is Precision@10 so low?",
         "There is only one relevant holding per query, so the best possible Precision@10 is 1/10 = 0.1. Our best "
         f"run reaches {f(HFT, 'P@10')}. It is low by design, not because the system is bad."),
        ("Why is MAP equal to MRR?",
         "With one relevant document, average precision of a query is 1 / rank of that document, which is exactly "
         "the reciprocal rank."),
        ("How did you draw the Precision-Recall graph?",
         "Mean Precision@k against mean Recall@k for k = 1 to 100. The classic 11-point interpolated curve would be "
         "a flat line with one relevant document, so this version is more informative. Precision is on a log scale."),
        ("What does nDCG@10 mean here?",
         "With binary relevance it is 1 / log2(rank + 1) if the holding is in the top 10 and 0 otherwise, averaged "
         "over queries. It rewards finding the answer near the top."),
        ("Is the hybrid's improvement statistically significant?",
         "Yes. A Wilcoxon signed-rank test on per-query reciprocal ranks against BM25 gives p < 0.001 for the "
         "three-way hybrid, the fine-tuned model and the fine-tuned hybrid. BM25 + base dense is not significantly "
         "different from BM25 (p = 0.46)."),
    ]),
    ("Results", [
        ("What are your main numbers?",
         f"Full pool, test: TF-IDF MAP {f('TF-IDF', 'MAP')}, BM25 {f('BM25', 'MAP')}, Sentence Transformer "
         f"{f('Dense', 'MAP')}, hybrid RRF {f(H3, 'MAP')} (Recall@100 {f(H3, 'R@100')} vs {f('BM25', 'R@100')} for "
         f"BM25). Best overall: BM25 + fine-tuned dense, MAP {f(HFT, 'MAP')}."),
        ("Why is the dense model worse than BM25?",
         "It was trained on general web and question-answer data, not court language. Legal matching often depends "
         "on exact rare terms (statute names, doctrines) that BM25 catches. Long passages are also cut at 256 word "
         "pieces."),
        ("If the dense model is weaker, why does fusion still help?",
         f"Because it fails on different queries. BM25 alone gets rank 1 in {err['bm25_top1_dense_not']} queries "
         f"where the dense model doesn't, and the dense model alone in {err['dense_top1_bm25_not']}. RRF combines "
         "both, so answers that one model buried get pulled into the top 100."),
        ("Why does the hybrid not help in the 5-choice setting?",
         f"There BM25 gets {f('BM25', 'accuracy', five)} and the hybrid {f(H3, 'accuracy', five)}. The wrong options "
         "were chosen by TF-IDF similarity, so all five already share words; extra lexical signal does not separate "
         "them."),
        ("How did fine-tuning work?",
         "Same MiniLM, one epoch over 45,000 training pairs (passage, correct holding) with one wrong option as a "
         "hard negative, multiple negatives ranking loss, batch 32, learning rate 2e-5, about 11 minutes on the "
         f"laptop GPU. MAP went from {f('Dense', 'MAP')} to {f(FT, 'MAP')}."),
        ("Is fine-tuning on CaseHOLD cheating?",
         "No. We trained only on the official train split and evaluated on the separate test split. Using training "
         "data is normal; we just report it as an extension beyond the brief's four models."),
        ("Why did the cross-encoder make things worse?",
         f"It was trained on MS MARCO web search: short questions and web passages. Our queries are ~140-word legal "
         f"passages. Out of its domain it reorders the top 20 badly (MAP {f(CE, 'MAP')})."),
        ("How do your results compare with published work?",
         "The CaseHOLD paper reports macro F1 of about 0.4 for a BiLSTM and about 0.6 for fine-tuned BERT on the "
         f"5-choice task. Our accuracies (BM25 {f('BM25', 'accuracy', five)}, fine-tuned MiniLM "
         f"{f(FT, 'accuracy', five)}) are on the LexGLUE split with a different metric, so they are not directly "
         "comparable, but they are in a sensible range."),
    ]),
    ("Errors, limits and demo", [
        ("Give an example where BM25 wins.",
         "Query 3185 is about money 'structuring' offences. That rare word appears in the right holding, so BM25 "
         "ranks it first; the dense model ranks it 129th."),
        ("And one where the dense model wins?",
         "Query 296: the passage and the holding share no keywords (relationships, cease contact vs restricting "
         "police officers' sexual conduct). BM25 ranks it 1,694th, the dense model first."),
        ("Which queries can't be solved at all?",
         "Query 2885: the passage is about a strip search but the citation supports a procedural point about one "
         "court panel being bound by another. Nothing in the text says that, so no ranker can find it."),
        ("What are the limitations?",
         "One relevant holding per query, so a legally valid but different holding counts as wrong; the pool is the "
         "test split's holdings, not a full case law database; small models and one epoch because everything ran "
         "on a laptop."),
        ("What would you do next?",
         "Use a legal encoder such as Legal-BERT, train the cross-encoder on CaseHOLD, and learn weights for each "
         "ranker inside RRF."),
        ("How do we run it?",
         "pip install -r requirements.txt, then bash run_all.sh reproduces every table and chart. The demo is "
         "streamlit run app/search_app.py."),
        ("How long does it take?",
         "BM25 and TF-IDF index the test pool in under a second and score all 3,600 queries in about half a "
         "second each. Dense encoding takes about 14 s on the laptop GPU, fine-tuning about 11 minutes, and the "
         "cross-encoder about 3 minutes."),
        ("Who did what?",
         "Fill in before the viva. Suggested split for presenting: Vaibhav slides 1 to 5 (problem, data, setup), "
         "Pranav slides 6 to 11 (method, tuning, main results), Ishaan slides 12 to 17 (5-choice, extensions, "
         "errors, demo, conclusion)."),
    ]),
]

doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Inches(0.9)
sec.top_margin = sec.bottom_margin = Inches(0.8)
st = doc.styles["Normal"]
st.font.name, st.font.size = "Calibri", Pt(10.5)

p = doc.add_paragraph()
r = p.add_run("Viva preparation: Legal Retrieval using CaseHOLD")
r.bold, r.font.size, r.font.color.rgb = True, Pt(18), NAVY
p = doc.add_paragraph()
r = p.add_run("Group 13 (Vaibhav Bansal, Pranav Joshi, Ishaan Manocha) · CSE4144 Information Retrieval. "
              "All numbers are from our test-split run (results/ folder).")
r.font.size, r.font.color.rgb = Pt(10), RGBColor(0x59, 0x59, 0x59)

n = 0
for section, items in QA:
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(10)
    h.paragraph_format.keep_with_next = True
    r = h.add_run(section)
    r.bold, r.font.size, r.font.color.rgb = True, Pt(13), TEAL
    for q, a in items:
        n += 1
        qp = doc.add_paragraph()
        qp.paragraph_format.space_after = Pt(1)
        qp.paragraph_format.keep_with_next = True
        r = qp.add_run(f"{n}. {q}")
        r.bold, r.font.color.rgb = True, NAVY
        ap = doc.add_paragraph(a)
        ap.paragraph_format.space_after = Pt(6)
        ap.paragraph_format.left_indent = Inches(0.25)

out = R("deliverables/IR_Group13_Viva_QA.docx")
doc.save(out)
print("saved", out, n, "questions")
